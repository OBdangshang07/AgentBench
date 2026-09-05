from __future__ import annotations

import ast
import base64
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from agentbench.catalog import (
    NCRE_OFFICE_PAPER02_SUITE_ID,
    NCRE_OFFICE_PAPER03_SUITE_ID,
    NCRE_OFFICE_SUITE_ID,
    build_catalog,
)
from agentbench.execution import Workspace
from agentbench.ncre_assets import blobs
from agentbench.service import EvaluationService, public_definition

REPO_ROOT = Path(__file__).resolve().parents[2]
NCRE_PAPERS = ["paper01", "paper02", "paper03"]
NCRE_SECTIONS = ["choice", "word", "excel", "ppt"]
NCRE_SLUGS = [f"ncre.office.{paper}.{section}" for paper in NCRE_PAPERS for section in NCRE_SECTIONS]
NCRE_METRIC_KEYS = {
    "judge_choice.py": [f"q{i:02d}" for i in range(1, 21)],
    "judge_word.py": [f"w{i}" for i in range(1, 10)],
    "judge_excel.py": [f"e{i}" for i in range(1, 10)],
    "judge_ppt.py": [f"p{i}" for i in range(1, 9)],
    "judge_choice_paper02.py": [f"q{i:02d}" for i in range(1, 21)],
    "judge_word_paper02.py": [f"w{i}" for i in range(1, 10)],
    "judge_excel_paper02.py": [f"e{i}" for i in range(1, 9)],
    "judge_ppt_paper02.py": [f"p{i}" for i in range(1, 8)],
    "judge_choice_paper03.py": [f"q{i:02d}" for i in range(1, 21)],
    "judge_word_paper03.py": [f"w{i}" for i in range(1, 10)],
    "judge_excel_paper03.py": [f"e{i}" for i in range(1, 9)],
    "judge_ppt_paper03.py": [f"p{i}" for i in range(1, 8)],
}
EXPECTED_POINTS = {
    "ncre.office.paper01.choice": ("choice", 20),
    "ncre.office.paper01.word": ("word", 30),
    "ncre.office.paper01.excel": ("excel", 30),
    "ncre.office.paper01.ppt": ("ppt", 20),
    "ncre.office.paper02.choice": ("choice", 20),
    "ncre.office.paper02.word": ("word", 30),
    "ncre.office.paper02.excel": ("excel", 30),
    "ncre.office.paper02.ppt": ("ppt", 20),
    "ncre.office.paper03.choice": ("choice", 20),
    "ncre.office.paper03.word": ("word", 30),
    "ncre.office.paper03.excel": ("excel", 30),
    "ncre.office.paper03.ppt": ("ppt", 20),
}
JUDGE_BY_SLUG = {
    "ncre.office.paper01.choice": ("judge_choice.py", 20),
    "ncre.office.paper01.word": ("judge_word.py", 30),
    "ncre.office.paper01.excel": ("judge_excel.py", 30),
    "ncre.office.paper01.ppt": ("judge_ppt.py", 20),
    "ncre.office.paper02.choice": ("judge_choice_paper02.py", 20),
    "ncre.office.paper02.word": ("judge_word_paper02.py", 30),
    "ncre.office.paper02.excel": ("judge_excel_paper02.py", 30),
    "ncre.office.paper02.ppt": ("judge_ppt_paper02.py", 20),
    "ncre.office.paper03.choice": ("judge_choice_paper03.py", 20),
    "ncre.office.paper03.word": ("judge_word_paper03.py", 30),
    "ncre.office.paper03.excel": ("judge_excel_paper03.py", 30),
    "ncre.office.paper03.ppt": ("judge_ppt_paper03.py", 20),
}
REFERENCE_DIRS = {
    "paper01": "reference-workspace",
    "paper02": "reference-workspace-paper02",
    "paper03": "reference-workspace-paper03",
}
NCRE_SUITES = [
    (NCRE_OFFICE_SUITE_ID, "paper01"),
    (NCRE_OFFICE_PAPER02_SUITE_ID, "paper02"),
    (NCRE_OFFICE_PAPER03_SUITE_ID, "paper03"),
]
# service.py _validate_definition 支持的 validator 白名单（保持同步）
SUPPORTED_VALIDATORS = {
    "exact_match",
    "contains",
    "regex",
    "json_schema",
    "file_exists",
    "file_content",
    "file_contains",
    "forbidden_paths",
    "command",
    "command_metrics",
    "ai_rubric",
}


def ncre_cases() -> dict[str, dict]:
    cases = build_catalog()
    return {
        case["slug"]: case for case in cases if case["category"] == "office-exam"
    }


def case_private_files(case: dict) -> dict[str, str]:
    """私有判分脚本存放于 command_metrics validator 的 config 层。"""
    command = next(
        item for item in case["validators"] if item["type"] == "command_metrics"
    )
    return command["config"]["private_files"]


def judge_answers(paper: str = "paper01") -> dict[str, str]:
    """从私有选择题判分脚本源码中提取标准答案表（不硬编码答案）。"""
    slug = f"ncre.office.{paper}.choice"
    private_files = case_private_files(ncre_cases()[slug])
    judge_name = next(iter(private_files))
    tree = ast.parse(private_files[judge_name])
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id == "ANSWERS":
                    return ast.literal_eval(node.value)
    raise AssertionError(f"{judge_name} 中未找到 ANSWERS 答案表")


def test_catalog_contains_212_cases_with_ncre_appended():
    cases = build_catalog()
    assert len(cases) == 212
    ncre = [case for case in cases if case["category"] == "office-exam"]
    assert len(ncre) == 12
    assert {case["slug"] for case in ncre} == set(NCRE_SLUGS)
    for case in ncre:
        assert case["version"] == "1.0.0"
        assert case["metadata"]["exam"] == "ncre-office"
        assert case["metadata"]["capability"] == "office-application"


def test_ncre_validators_weight_total_and_whitelist():
    for slug, case in ncre_cases().items():
        validators = case["validators"]
        assert sum(item["weight"] for item in validators) == 100, slug
        types = {item["type"] for item in validators}
        assert types <= SUPPORTED_VALIDATORS, slug
        assert types == {"command_metrics", "file_exists", "forbidden_paths"}, slug
        # service 层定义校验应直接通过（权重和与类型白名单双重确认）
        EvaluationService._validate_definition(case)


def test_ncre_metadata_exam_fields():
    for slug, case in ncre_cases().items():
        metadata = case["metadata"]
        for key in ("exam", "exam_section", "exam_points", "source"):
            assert key in metadata, f"{slug} metadata 缺少 {key}"
        section, points = EXPECTED_POINTS[slug]
        assert metadata["exam_section"] == section
        assert metadata["exam_points"] == points
        assert metadata["source"]


def test_ncre_command_metrics_keys_match_judges():
    for slug, (judge, points) in JUDGE_BY_SLUG.items():
        case = ncre_cases()[slug]
        command = next(
            item for item in case["validators"] if item["type"] == "command_metrics"
        )
        assert command["weight"] == 90
        assert judge in command["config"]["command"]
        assert command["config"]["private_files"]
        keys = [metric["key"] for metric in command["config"]["metrics"]]
        assert keys == NCRE_METRIC_KEYS[judge]
        assert sum(metric["weight"] for metric in command["config"]["metrics"]) == points


@pytest.mark.parametrize("paper", NCRE_PAPERS)
def test_choice_instruction_does_not_leak_answers(paper):
    instruction = ncre_cases()[f"ncre.office.{paper}.choice"]["instruction"]
    answers = judge_answers(paper)
    assert len(answers) == 20
    # instruction 固定携带格式示例行（{"q01": "A", "q02": "B", ...}），
    # 先剔除再判定，避免与真实答案巧合重叠造成误报
    example_line = '例如 {"q01": "A", "q02": "B", ...}。'
    assert example_line in instruction
    body = instruction.replace(example_line, "")
    # 答案映射特征不得出现在面向被测 Agent 的 instruction 中
    for qid, letter in answers.items():
        assert f'"{qid}": "{letter}"' not in body, f"{paper} {qid} 答案泄漏"
        assert f'"{qid}":"{letter}"' not in body, f"{paper} {qid} 答案泄漏"
    assert json.dumps(answers) not in instruction
    assert json.dumps(answers, separators=(",", ":")) not in instruction
    assert "".join(answers[key] for key in sorted(answers)) not in instruction
    assert "ANSWERS" not in instruction
    assert "judge_choice" not in instruction
    # 题干与选项必须完整给出（否则无法作答）
    for qid in answers:
        assert qid in instruction


def test_private_judge_scripts_compile():
    for slug, case in ncre_cases().items():
        assert len(case_private_files(case)) == 1, slug
        for name, source in case_private_files(case).items():
            compile(source, name, "exec")
            assert set(NCRE_METRIC_KEYS[name])


def test_private_judges_are_hidden_from_public_definition():
    for slug, case in ncre_cases().items():
        public = public_definition(case)
        assert "private_files" not in public, slug
        command = next(
            item for item in public["validators"] if item["type"] == "command_metrics"
        )
        config = command["config"]
        assert "private_files" not in config, slug
        assert config.get("private") is True
        assert config["command"] == "<AgentBench private validator>"
        assert "judge_" not in json.dumps(public, ensure_ascii=False)


def test_workspace_seed_supports_plain_text_and_base64(tmp_path):
    workspace = Workspace(tmp_path)
    raw_bytes = b"\x89PNG\r\n\x1a\n\x00\x01\x02\xff\xfe"
    workspace.write_file("plain.txt", "纯文本内容 unchanged")
    workspace.write_file("blob.bin", "base64:" + base64.b64encode(raw_bytes).decode("ascii"))
    assert (tmp_path / "plain.txt").read_bytes() == "纯文本内容 unchanged".encode()
    assert (tmp_path / "blob.bin").read_bytes() == raw_bytes

    # 目录级 seed：真实 NCRE initial_files（base64 与纯文本混合）
    word_case = ncre_cases()["ncre.office.paper01.word"]
    mixed_root = tmp_path / "mixed"
    mixed = Workspace(mixed_root)
    mixed.seed(word_case["initial_files"])
    assert (mixed_root / "通讯录.csv").read_text(encoding="utf-8") == blobs.CONTACTS_CSV
    docx_bytes = (mixed_root / "Word.docx").read_bytes()
    assert docx_bytes == base64.b64decode(blobs.WORD_DOCX_B64)
    assert zipfile.is_zipfile(mixed_root / "Word.docx")


@pytest.fixture(scope="module")
def reference_workspaces(tmp_path_factory) -> dict[str, Path]:
    """逐卷运行 _ncre_dev/make_reference.py 生成满分标准产物并复制到临时目录。"""
    workspaces: dict[str, Path] = {}
    for paper, dirname in REFERENCE_DIRS.items():
        result = subprocess.run(
            [
                sys.executable,
                str(REPO_ROOT / "_ncre_dev" / "make_reference.py"),
                paper,
            ],
            cwd=REPO_ROOT,
            capture_output=True,
            text=True,
            timeout=600,
            check=False,
        )
        assert result.returncode == 0, f"{paper}: {result.stderr}"
        source = REPO_ROOT / "_ncre_dev" / dirname
        assert source.exists(), f"{paper} 标准产物目录缺失: {source}"
        workspace = tmp_path_factory.mktemp(f"ncre-{paper}")
        for item in source.iterdir():
            shutil.copy2(item, workspace / item.name)
        workspaces[paper] = workspace
    return workspaces


def run_judge(case: dict, workspace: Path, tmp_path: Path) -> dict:
    private_root = tmp_path / "private"
    private_root.mkdir(exist_ok=True)
    payload = {}
    for name, source in case_private_files(case).items():
        script = private_root / name
        script.write_text(source, encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(script)],
            cwd=workspace,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=300,
            check=False,
        )
        assert result.returncode == 0, f"{name}: {result.stderr}"
        line = next(
            (
                raw
                for raw in result.stdout.splitlines()
                if raw.startswith("AGENTBENCH_METRICS=")
            ),
            None,
        )
        assert line is not None, f"{name} 未输出协议行: {result.stdout}"
        payload[name] = json.loads(line.removeprefix("AGENTBENCH_METRICS="))
    return payload


def test_ppt_judge_accepts_visual_titles_hierarchy_and_standard_custom_shows(tmp_path):
    """自定义布局不应因未使用标题/正文占位符而丢失内容分。

    OOXML 自定义放映可以使用 p:custShow，且 Relationship 属性没有
    固定顺序；判分器需按 XML 语义解析，不能依赖特定生成器格式。
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    titles = [
        "Microsoft Office图书策划案",
        "推荐作者简介",
        "Office 2010的十大优势",
        "新版图书读者定位",
        "PowerPoint 2010创新的功能体验",
        "2012年同类图书销量统计",
        "新版图书创作流程示意",
    ]
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = titles[0]
    for title in titles[1:]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.6))
        box.text_frame.text = title

    def add_text(slide, text, top, size):
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(8), Inches(0.5))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)

    page2 = presentation.slides[1]
    add_text(page2, "刘雅汶", 1.3, 24)
    add_text(page2, "Contoso公司技术经理", 1.9, 12)
    add_text(page2, "主要代表作品", 2.7, 20)
    add_text(page2, "《Microsoft Office整合应用精要》", 3.3, 12)
    page4 = presentation.slides[3]
    for index, text in enumerate(("信息工作者", "学生和教师", "办公应用技能培训班", "大专院校教材")):
        add_text(page4, text, 1.2 + index * 0.8, 18)

    source = tmp_path / "source.pptx"
    presentation.save(source)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    target = workspace / "PowerPoint.pptx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "w") as rewritten:
        presentation_xml = original.read("ppt/presentation.xml").decode("utf-8")
        rels_xml = original.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
        slide_ids = re.findall(r'<p:sldId\b[^>]*r:id="([^"]+)"', presentation_xml)
        assert len(slide_ids) == 7
        custom_shows = (
            '<p:custShowLst xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<p:custShow name="放映方案1" id="1"><p:sldLst>'
            + "".join(f'<p:sld r:id="{slide_ids[i - 1]}"/>' for i in (1, 2, 4, 7))
            + '</p:sldLst></p:custShow><p:custShow name="放映方案2" id="2"><p:sldLst>'
            + "".join(f'<p:sld r:id="{slide_ids[i - 1]}"/>' for i in (1, 2, 3, 5, 6))
            + "</p:sldLst></p:custShow></p:custShowLst>"
        )
        presentation_xml = presentation_xml.replace(
            "</p:presentation>", custom_shows + "</p:presentation>"
        )
        rels_xml = re.sub(
            r'<Relationship Id="([^"]+)" Type="([^"]+)" Target="([^"]+)"',
            r'<Relationship Type="\2" Target="\3" Id="\1"',
            rels_xml,
        )
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                data = presentation_xml.encode("utf-8")
            elif item.filename == "ppt/_rels/presentation.xml.rels":
                data = rels_xml.encode("utf-8")
            rewritten.writestr(item, data)

    case = ncre_cases()["ncre.office.paper01.ppt"]
    result = run_judge(case, workspace, tmp_path)["judge_ppt.py"]["metrics"]
    assert result["p2"] == 100.0
    assert result["p3"] == 100.0
    assert result["p8"] == 100.0


@pytest.mark.parametrize("slug", NCRE_SLUGS)
def test_private_judges_give_full_marks_on_reference_artifacts(
    slug, reference_workspaces, tmp_path
):
    case = ncre_cases()[slug]
    paper = slug.split(".")[2]
    judge_name = next(iter(case_private_files(case)))
    payloads = run_judge(case, reference_workspaces[paper], tmp_path)
    metrics = payloads[judge_name]["metrics"]
    expected_keys = NCRE_METRIC_KEYS[judge_name]
    assert set(metrics) == set(expected_keys)
    bad = {key: metrics[key] for key in expected_keys if metrics[key] != 100.0}
    assert not bad, f"{judge_name} 标准产物未满分: {bad}\n{payloads[judge_name].get('evidence')}"


@pytest.mark.parametrize("paper", NCRE_PAPERS)
def test_choice_judge_empty_workspace_emits_zero_metrics(paper, tmp_path):
    case = ncre_cases()[f"ncre.office.{paper}.choice"]
    judge_name = next(iter(case_private_files(case)))
    empty = tmp_path / "empty"
    empty.mkdir()
    payloads = run_judge(case, empty, tmp_path)
    metrics = payloads[judge_name]["metrics"]
    assert set(metrics) == {f"q{i:02d}" for i in range(1, 21)}
    assert all(value == 0.0 for value in metrics.values())


@pytest.mark.parametrize("suite_id,paper", NCRE_SUITES)
def test_ncre_suite_registered_with_four_members(settings, suite_id, paper):
    service = EvaluationService(settings)
    try:
        suite = service.get_suite(suite_id)
        members = {item["slug"] for item in suite["cases"]}
        assert members == {f"ncre.office.{paper}.{section}" for section in NCRE_SECTIONS}
    finally:
        service.close()
